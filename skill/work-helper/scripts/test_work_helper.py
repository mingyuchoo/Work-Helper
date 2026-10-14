"""임시 엑셀로 CLI의 업무 결과와 실패 시 동작을 검증한다."""

import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook


CLI = Path(__file__).resolve().with_name("work_helper.py")
REPO = CLI.parents[3]


class WorkHelperTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def workbook(self, name, sheet, rows):
        path = self.root / name
        wb = Workbook()
        wb.active.title = sheet
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        wb.close()
        return path

    def run_cli(self, *args, success=True, cli=CLI):
        result = subprocess.run(
            [sys.executable, str(cli), *map(str, args)], cwd=self.root,
            capture_output=True, text=True, encoding="utf-8",
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )
        self.assertEqual(result.returncode, 0 if success else 1, result.stderr or result.stdout)
        data = json.loads(result.stdout)
        self.assertEqual(data["ok"], success)
        return data

    def test_merge_values_and_sources(self):
        for name, value in (("실적_a.xlsx", 10), ("실적_b.xlsx", 20)):
            self.workbook(name, "실적", [["금액"], [value]])
        args = ("merge", "--folder", self.root, "--prefix", "실적_", "--sheets", "실적")
        self.run_cli(*args, "--headers", "1,2", success=False)
        self.assertEqual(len(list(self.root.glob("*.xlsx"))), 2)
        data = self.run_cli(*args)
        self.assertEqual(data["total_rows"], 2)
        wb = load_workbook(data["out_path"])
        try:
            rows = list(wb.active.values)
            self.assertEqual(rows[0], ("출처파일", "출처폴더", "금액"))
            self.assertEqual([r[2] for r in rows[1:]], [10, 20])
            self.assertEqual([r[0] for r in rows[1:]], ["실적_a", "실적_b"])
        finally:
            wb.close()

    def test_recipients_and_saved_text(self):
        roster = self.workbook("주소록.xlsx", "주소록", [
            ["사업부", "팀", "역할", "이름", "이메일"],
            ["제품", "개발팀", "팀장", "가", "a@example.com"],
            ["제품", "개발팀", "팀원", "나", "b@example.com"],
            ["영업", "영업팀", "팀원", "다", "c@example.com"],
        ])
        meta = self.run_cli("roster", "--file", roster)
        self.assertEqual(meta["roster_count"], 3)
        self.assertNotIn("_state", meta)
        args = ("recipients", "--file", roster, "--teams", "개발팀", "--to-roles", "팀장")
        data = self.run_cli(*args, "--cc-roles", "팀장,팀원", "--save-txt")
        self.assertEqual(data["to_join"], "a@example.com")
        # 기존 백엔드는 일반 역할을 주소록 전체에서 선택한다.
        self.assertEqual(data["cc_join"], "b@example.com; c@example.com")
        self.assertIn("b@example.com", Path(data["out_path"]).read_text(encoding="utf-8-sig"))
        self.run_cli(*args, "--cc-roles", "없는역할", success=False)
        self.run_cli("roster", "--file", self.root / "missing.xlsx", success=False)

    def test_folder_preview_create_and_invalid_column(self):
        excel = self.workbook("폴더.xlsx", "목록", [
            ["사업부", "팀"], ["제품", "개발"], ["제품", "개발"], ["제품", "품질"],
        ])
        output = self.root / "output"
        meta = self.run_cli("folder-info", "--file", excel)
        self.assertEqual(meta["sheets"][0]["headers"], ["사업부", "팀"])
        args = ("folders", "--file", excel, "--sheet", "목록", "--upper-col", "사업부")
        data = self.run_cli(*args, "--lower-col", "팀", "--output-dir", output)
        self.assertEqual(data["tree"][0]["children"], [{"name": "01.개발"}, {"name": "02.품질"}])
        self.assertFalse(output.exists())
        self.run_cli(*args, "--lower-col", "없는열", "--create", "--output-dir", output, success=False)
        self.assertFalse(output.exists())
        self.run_cli(*args, "--create", success=False)
        data = self.run_cli(*args, "--lower-col", "팀", "--create", "--output-dir", output)
        self.assertEqual(data["created"], 2)
        self.assertTrue((output / "01.제품" / "02.품질").is_dir())
        data = self.run_cli(*args, "--lower-col", "팀", "--create", "--output-dir", output)
        self.assertEqual((data["created"], data["existed"]), (0, 2))

    def test_copied_cli_uses_explicit_repository(self):
        installed = self.root / "installed" / "work-helper" / "scripts" / CLI.name
        installed.parent.mkdir(parents=True)
        shutil.copyfile(CLI, installed)
        excel = self.workbook("columns.xlsx", "목록", [["상위", "하위"], ["A", "B"]])
        self.run_cli("folder-info", "--file", excel, cli=installed, success=False)
        data = self.run_cli("--repo-root", REPO, "folder-info", "--file", excel, cli=installed)
        self.assertEqual(data["sheets"][0]["name"], "목록")


if __name__ == "__main__":
    unittest.main()
