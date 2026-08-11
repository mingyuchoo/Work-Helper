# -*- coding: utf-8 -*-
"""메일 수신자/참조자 생성 기능 (메일주소_생성.py 의 핵심 로직을 함수로 재사용).

원본의 input()/파일선택창/print 대신, 웹 UI 에서 넘어온 인자로 동작한다.
동작 2단계:
  1) analyze(path)  : 주소록을 읽어 머리글/열 자동 인식 → 역할·팀 목록 반환 + 파싱 상태 보관
  2) compute(state, teams_text, to_roles, cc_roles) : 수신/참조 이메일 계산
"""

import re
import os
from datetime import datetime

import openpyxl


# ──────────────────────────────────────────────────────────────
# 공통 유틸
# ──────────────────────────────────────────────────────────────
def norm(s):
    """비교용 정규화: 문자열로 바꾼 뒤 모든 공백 제거."""
    return re.sub(r"\s+", "", "" if s is None else str(s))


def cellval(rec, c):
    """레코드에서 c번째 열 값을 문자열(앞뒤 공백 제거)로 꺼낸다. 열이 없으면 ''."""
    if c is None:
        return ""
    v = rec[c] if c < len(rec) else None
    return "" if v is None else str(v).strip()


_BLANK_MARKS = {"", "-", "--", "n/a", "na", "없음"}


def orgval(rec, c):
    """조직/팀 열 값. '-' 등 빈칸 표기는 ''(비어 있음)으로 처리한다."""
    v = cellval(rec, c)
    return "" if v.lower() in _BLANK_MARKS else v


COL_KEYWORDS = {
    "team":  ["팀", "부서", "소속조직"],
    "upper": ["사업부", "부문", "본부", "상위조직"],
    "name":  ["이름", "성명", "담당자명", "성함"],
    "email": ["이메일", "메일", "email", "e-mail", "mail"],
    "role":  ["역할", "구분", "직책", "직위", "role", "구성원구분"],
}


def read_matrix(path):
    """엑셀 첫 시트(또는 값이 있는 첫 시트)를 2차원 값 배열로 읽는다.
    병합 셀은 대표값(왼쪽 위)으로 채운다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    ws = None
    for cand in wb.worksheets:
        if cand.max_row and cand.max_column:
            ws = cand
            break
    if ws is None:
        ws = wb.active

    max_row, max_col = ws.max_row, ws.max_column
    matrix = [[ws.cell(r, c).value for c in range(1, max_col + 1)]
              for r in range(1, max_row + 1)]
    for rng in ws.merged_cells.ranges:
        top = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                matrix[r - 1][c - 1] = top
    wb.close()
    while matrix and all(v is None or str(v).strip() == "" for v in matrix[-1]):
        matrix.pop()
    return matrix


def score_header(row):
    cells = [norm(v).lower() for v in row]
    hit = 0
    for keys in COL_KEYWORDS.values():
        if any(any(k in cell for k in keys) for cell in cells):
            hit += 1
    return hit


def find_header_row(matrix, scan=15):
    best_idx, best_score = 0, -1
    for i in range(min(scan, len(matrix))):
        s = score_header(matrix[i])
        if s > best_score:
            best_idx, best_score = i, s
    return best_idx, best_score


def detect_column(header, kind):
    keys = [k.lower() for k in COL_KEYWORDS[kind]]
    best_idx, best_q = None, 0
    for idx, h in enumerate(header):
        cell = norm(h).lower()
        if not cell:
            continue
        q = 0
        for k in keys:
            if cell == k:
                q = max(q, 3)
            elif cell.startswith(k) or cell.endswith(k):
                q = max(q, 2)
            elif k in cell:
                q = max(q, 1)
        if q > best_q:
            best_idx, best_q = idx, q
    return best_idx


def parse_list(raw):
    return [s.strip() for s in str(raw).split(",") if s.strip()]


# '전체'를 뜻하는 입력들
ALL_TOKENS = {"전체", "전부", "모두", "전사", "all", "*"}


def match_teams(records, team_col, upper_cols, wanted):
    """입력 토큰을 '팀명 + 상위 조직(사업부/Level·담당)'과 대조해 대상 팀 집합을 만든다.
    - '전체'(전부/모두/전사/all/*): 주소록의 모든 팀
    - 사업부명 / 담당명(상위 조직): 그 조직에 속한 모든 팀
    - 팀명: 해당 팀
    반환: (매칭된 실제 팀명 집합, 못 찾은 입력 목록)"""
    # 팀 목록 (정규화키 -> 표시용 원본)
    all_teams = {}
    for rec in records:
        t = orgval(rec, team_col)
        if t:
            all_teams.setdefault(norm(t), t)

    matched, missing = set(), []
    for w in wanted:
        key = norm(w)
        if not key:
            continue
        # 0) 전체 지정
        if key.lower() in ALL_TOKENS:
            matched.update(all_teams.values())
            continue

        found = set()
        # 1) 팀명 매칭 (정확 → 부분)
        if key in all_teams:
            found.add(all_teams[key])
        else:
            found.update(orig for k, orig in all_teams.items()
                         if key in k or k in key)
        # 2) 상위 조직(사업부/Level·담당) 매칭 → 그 조직에 속한 모든 팀
        for rec in records:
            t = orgval(rec, team_col)
            if not t:
                continue
            for c in upper_cols:
                v = orgval(rec, c)
                if v and (norm(v) == key or key in norm(v)):
                    found.add(t)
                    break

        if found:
            matched.update(found)
        else:
            missing.append(w)
    return matched, missing


EMAIL_RX = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")


def compute_upper_cols(team_col, exclude):
    if team_col is None:
        return []
    return [c for c in range(0, team_col) if c not in exclude]


def build_candidates(records, team_col, upper_cols, matched_teams):
    team_keys = {norm(t) for t in matched_teams}

    team_upper = {}
    for rec in records:
        t = orgval(rec, team_col)
        if not (t and norm(t) in team_keys):
            continue
        prof = team_upper.setdefault(norm(t), {})
        for c in upper_cols:
            v = orgval(rec, c)
            if v:
                prof.setdefault(c, set()).add(norm(v))

    no_upper = (sorted({t for t in matched_teams if not team_upper.get(norm(t))})
                if upper_cols else [])

    cands = []
    for rec in records:
        t = orgval(rec, team_col)
        if t and norm(t) in team_keys:
            cands.append(rec)
            continue
        if t or not upper_cols:
            continue
        rvals = {c: norm(orgval(rec, c)) for c in upper_cols if orgval(rec, c)}
        if not rvals:
            continue
        for prof in team_upper.values():
            if all(c in prof and rvals[c] in prof[c] for c in rvals):
                cands.append(rec)
                break
    return cands, no_upper


def upper_label(rec, upper_cols):
    vals = [orgval(rec, c) for c in upper_cols if orgval(rec, c)]
    return "[" + " > ".join(vals) + "]" if vals else "(상위)"


POSITION_ROLE_EXACT = {"팀qa", "담당", "담당임원"}
POSITION_ROLE_LABEL = "팀QA·팀장·단장/실장·담당·센터장·사업부장 등 직책"


def is_position_role(role):
    r = norm(role).lower()
    if not r:
        return False
    if r in POSITION_ROLE_EXACT:
        return True
    return r.endswith("장")


def collect(records, cand_ids, team_col, upper_cols, role_col, name_col, email_col, roles):
    role_set = set(roles)
    out, seen = [], set()
    warnings = []
    for rec in records:
        role = cellval(rec, role_col)
        if role not in role_set:
            continue
        if is_position_role(role) and id(rec) not in cand_ids:
            continue
        m = EMAIL_RX.search(cellval(rec, email_col))
        email = m.group(0) if m else cellval(rec, email_col)
        name = cellval(rec, name_col)
        team = orgval(rec, team_col) or upper_label(rec, upper_cols)
        if not email:
            warnings.append(f"이메일 없음 → 건너뜀: {name or '(이름없음)'} / {team} / {role}")
            continue
        key = email.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append({"name": name, "team": team, "role": role, "email": email})
    return out, warnings


# ──────────────────────────────────────────────────────────────
# 웹 UI 용 상위 함수
# ──────────────────────────────────────────────────────────────
def analyze(path):
    """주소록 파일을 읽어 UI 에 필요한 정보를 반환한다.
    반환 dict 의 '_state' 는 compute() 에 그대로 넘긴다 (파이썬 측 보관용)."""
    matrix = read_matrix(path)
    if not matrix:
        return {"ok": False, "error": "주소록 내용이 비어 있습니다."}

    hidx, _ = find_header_row(matrix)
    header = matrix[hidx]

    team_col = detect_column(header, "team")
    role_col = detect_column(header, "role")
    name_col = detect_column(header, "name")
    email_col = detect_column(header, "email")

    if team_col is None or role_col is None or email_col is None:
        missing = []
        if team_col is None:
            missing.append("팀/부서")
        if role_col is None:
            missing.append("역할")
        if email_col is None:
            missing.append("이메일")
        return {"ok": False,
                "error": "필수 열을 찾지 못했습니다: " + ", ".join(missing)
                         + " (주소록 머리글을 확인하세요)"}

    upper_cols = compute_upper_cols(team_col, {name_col, role_col, email_col})
    records = matrix[hidx + 1:]

    roles = sorted({cellval(r, role_col) for r in records if cellval(r, role_col)})
    teams = sorted({orgval(r, team_col) for r in records if orgval(r, team_col)})

    def _hname(c):
        return "" if c is None or header[c] is None else str(header[c]).strip()

    state = {
        "path": path,
        "header": header,
        "records": records,
        "team_col": team_col,
        "role_col": role_col,
        "name_col": name_col,
        "email_col": email_col,
        "upper_cols": upper_cols,
    }
    meta = {
        "ok": True,
        "file_name": os.path.basename(path),
        "roster_count": len(records),
        "roles": roles,
        "teams": teams,
        "columns": {
            "team": _hname(team_col),
            "role": _hname(role_col),
            "name": _hname(name_col),
            "email": _hname(email_col),
            "upper": [_hname(c) for c in upper_cols],
        },
    }
    return {"meta": meta, "_state": state}


def compute(state, teams_text, to_roles, cc_roles):
    """수신/참조 이메일을 계산한다."""
    records = state["records"]
    team_col = state["team_col"]
    role_col = state["role_col"]
    name_col = state["name_col"]
    email_col = state["email_col"]
    upper_cols = state["upper_cols"]

    wanted = parse_list(teams_text)
    matched, missing = match_teams(records, team_col, upper_cols, wanted)
    if not matched:
        teams_all = sorted({orgval(r, team_col) for r in records if orgval(r, team_col)})
        return {"ok": False,
                "error": "대상 팀에 해당하는 사람이 없습니다. 팀명을 확인하세요.",
                "missing": wanted, "teams_all": teams_all}

    cands, no_upper = build_candidates(records, team_col, upper_cols, matched)
    cand_ids = {id(r) for r in cands}

    to_rows, w1 = collect(records, cand_ids, team_col, upper_cols,
                          role_col, name_col, email_col, to_roles)
    cc_rows, w2 = collect(records, cand_ids, team_col, upper_cols,
                          role_col, name_col, email_col, cc_roles)

    # 수신자에 이미 들어간 사람은 참조자에서 제외
    to_emails = {r["email"].lower() for r in to_rows}
    cc_rows = [r for r in cc_rows if r["email"].lower() not in to_emails]

    to_join = "; ".join(r["email"] for r in to_rows)
    cc_join = "; ".join(r["email"] for r in cc_rows)

    return {
        "ok": True,
        "matched_teams": sorted(matched),
        "missing": missing,
        "no_upper": no_upper,
        "warnings": w1 + w2,
        "to": to_rows,
        "cc": cc_rows,
        "to_join": to_join,
        "cc_join": cc_join,
    }


def save_txt(state, result, to_roles, cc_roles):
    """수신/참조 결과를 원본 주소록 폴더에 txt 로 저장하고 경로를 반환한다."""
    path = state["path"]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = os.path.dirname(os.path.abspath(path))
    out_path = os.path.join(out_dir, f"메일주소_{ts}.txt")

    lines = []
    lines.append("대상 팀: " + ", ".join(result.get("matched_teams", [])))
    lines.append("수신자 역할: " + (", ".join(to_roles) or "(없음)"))
    lines.append("참조자 역할: " + (", ".join(cc_roles) or "(없음)"))
    lines.append("")
    lines.append("[수신자(To) 이메일]")
    lines.append(result.get("to_join") or "(없음)")
    lines.append("")
    lines.append("[참조자(CC) 이메일]")
    lines.append(result.get("cc_join") or "(없음)")
    lines.append("")
    lines.append("-" * 40)
    lines.append("[상세 목록]")
    for label, rows in (("수신자(To)", result.get("to", [])),
                        ("참조자(CC)", result.get("cc", []))):
        lines.append(f"■ {label} ({len(rows)}명)")
        for r in rows:
            lines.append(f"  {r['name'] or '(이름없음)'} | {r['team']} | {r['role']} | {r['email']}")
        lines.append("")
    content = "\n".join(lines)

    with open(out_path, "w", encoding="utf-8-sig") as f:
        f.write(content)
    return out_path
