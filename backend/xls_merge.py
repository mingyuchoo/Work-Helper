# -*- coding: utf-8 -*-
"""엑셀 시트 취합 기능 (시트_취합.py 의 핵심 로직을 함수로 재사용).

동일 접두어로 시작하는 엑셀들에서 동일명 시트를 세로로 이어붙여 한 파일로 만든다.
원본의 input()/폴더선택창/print 대신, 웹 UI 인자로 동작한다.
"""

import os
import re
import glob
from datetime import datetime, date

import openpyxl
from openpyxl.utils import get_column_letter


def find_files(folder, prefix, recursive):
    pattern = "**/*.xlsx" if recursive else "*.xlsx"
    found = []
    for path in glob.glob(os.path.join(folder, pattern), recursive=recursive):
        name = os.path.basename(path)
        if name.startswith("~$"):
            continue
        if name.startswith(prefix):
            found.append(path)
    return sorted(found)


def resolve_sheet(wb, sheet_name):
    if sheet_name in wb.sheetnames:
        return sheet_name
    target = sheet_name.strip()
    for s in wb.sheetnames:
        if s.strip() == target:
            return s
    return None


def norm_row(row):
    cells = ["" if v is None else str(v).strip() for v in row]
    while cells and cells[-1] == "":
        cells.pop()
    return cells


def is_empty_row(row):
    return all(v is None or str(v).strip() == "" for v in row)


_DATE_PATTERNS = [
    (re.compile(r'^(\d{4})\s*[-./년]\s*(\d{1,2})\s*[-./월]\s*(\d{1,2})\s*일?\.?$'), 4),
    (re.compile(r'^(\d{2})\s*[-/년]\s*(\d{1,2})\s*[-/월]\s*(\d{1,2})\s*일?$'), 2),
    (re.compile(r'^(\d{4})(\d{2})(\d{2})$'), 4),
]


def coerce_date(value):
    if not isinstance(value, str):
        return value
    s = value.strip()
    if not s:
        return value
    for rx, ylen in _DATE_PATTERNS:
        m = rx.match(s)
        if m:
            y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if ylen == 2:
                y += 2000
            try:
                return datetime(y, mo, d)
            except ValueError:
                return value
    return value


def read_sheet_matrix(ws):
    max_row, max_col = ws.max_row, ws.max_column
    matrix = [[ws.cell(r, c).value for c in range(1, max_col + 1)]
              for r in range(1, max_row + 1)]
    for rng in ws.merged_cells.ranges:
        top_val = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                matrix[r - 1][c - 1] = top_val
    while matrix and all(v is None or str(v).strip() == "" for v in matrix[-1]):
        matrix.pop()
    return matrix


INVALID_TITLE = set(r':\/?*[]')


def safe_title(name, used):
    t = "".join("_" if ch in INVALID_TITLE else ch for ch in name.strip())
    t = (t[:26] + "_취합")[:31] if len(t) > 26 else t + "_취합"
    base, i = t[:31], 2
    while t in used:
        suffix = f"_{i}"
        t = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(t)
    return t


def aggregate_sheet(files, folder, sheet_name, header_row):
    combined_rows = []
    header = None
    n_cols = 0
    ok_count = 0
    skipped = []
    is_first = True

    for path in files:
        rel = os.path.relpath(path, folder)
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
        except Exception as e:
            skipped.append((rel, f"열기 실패: {e}"))
            continue

        real_sheet = resolve_sheet(wb, sheet_name)
        if real_sheet is None:
            skipped.append((rel, f"'{sheet_name}' 시트 없음"))
            wb.close()
            continue

        matrix = read_sheet_matrix(wb[real_sheet])
        wb.close()
        if not matrix:
            skipped.append((rel, "시트 내용이 비어 있음"))
            continue

        n_cols = max(n_cols, max(len(r) for r in matrix))
        file_base = os.path.splitext(os.path.basename(path))[0]
        src_folder = os.path.dirname(rel) or "."

        if is_first:
            if 1 <= header_row <= len(matrix):
                header = matrix[header_row - 1]
                data_rows = matrix[header_row:]
            else:
                header = None
                data_rows = matrix
            is_first = False
        else:
            data_rows = matrix[header_row:] if header is not None else matrix

        header_key = norm_row(header) if header is not None else None
        for row in data_rows:
            if is_empty_row(row):
                continue
            if header_key is not None and norm_row(row) == header_key:
                continue
            combined_rows.append((file_base, src_folder, row))

        ok_count += 1

    return header, combined_rows, n_cols, ok_count, skipped


def write_sheet(out_ws, header, combined_rows, n_cols):
    front = ["출처파일", "출처폴더"]
    if header is not None:
        out_ws.append(front + [h if h is not None else "" for h in header])
    else:
        out_ws.append(front + [get_column_letter(i + 1) for i in range(n_cols)])

    for file_base, src_folder, row in combined_rows:
        padded = list(row) + [None] * (n_cols - len(row))
        padded = [coerce_date(v) for v in padded]
        out_ws.append([file_base, src_folder] + padded)

    for row_cells in out_ws.iter_rows(min_row=2):
        for cell in row_cells:
            if isinstance(cell.value, (datetime, date)):
                cell.number_format = "yyyy-mm-dd"

    bold = openpyxl.styles.Font(bold=True)
    for cell in out_ws[1]:
        cell.font = bold
    out_ws.freeze_panes = "A2"
    out_ws.column_dimensions["A"].width = 30
    out_ws.column_dimensions["B"].width = 20


def parse_sheets(raw):
    return [s.strip() for s in str(raw).split(",") if s.strip()]


def resolve_header_rows(raw, n_sheets):
    parts = [p.strip() for p in str(raw).split(",") if p.strip()]
    try:
        nums = [int(p) for p in parts]
    except ValueError:
        nums = [1]
    if len(nums) == 1:
        return nums * n_sheets
    if len(nums) == n_sheets:
        return nums
    return [nums[0]] * n_sheets


def _fmt(v):
    """미리보기용 셀 표시 문자열."""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v)


# ── 웹 UI 용 상위 함수 ──────────────────────────────────────────
def run(folder, prefix, sheets_text, headers_text, recursive):
    """폴더 안 엑셀을 취합해 저장하고, 결과 요약/로그/미리보기를 반환한다."""
    if not os.path.isdir(folder):
        return {"ok": False, "error": f"폴더를 찾을 수 없습니다: {folder}"}

    sheets = parse_sheets(sheets_text)
    if not sheets:
        return {"ok": False, "error": "취합할 시트명이 없습니다."}

    header_rows = resolve_header_rows(headers_text, len(sheets))

    files = find_files(folder, prefix, recursive)
    if not files:
        return {"ok": False,
                "error": f"'{prefix}' 로 시작하는 xlsx 파일을 찾지 못했습니다."}

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)
    used_titles = set()
    total_rows = 0
    made = 0

    sheet_results = []
    all_log = []
    preview = None

    # 취합 대상/제외 파일을 로그로 남긴다.
    file_notes = {}  # rel -> (note, ok)

    for sheet_name, header_row in zip(sheets, header_rows):
        header, rows, n_cols, ok_count, skipped = aggregate_sheet(
            files, folder, sheet_name, header_row)

        for rel, reason in skipped:
            prev = file_notes.get(rel)
            note = f"[{sheet_name}] {reason}"
            if prev is None:
                file_notes[rel] = (note, False)

        sheet_results.append({
            "name": sheet_name,
            "header_row": header_row,
            "ok_count": ok_count,
            "rows": len(rows),
            "skipped": [{"file": rel, "reason": reason} for rel, reason in skipped],
        })

        if ok_count == 0:
            continue

        out_ws = out_wb.create_sheet(title=safe_title(sheet_name, used_titles))
        write_sheet(out_ws, header, rows, n_cols)
        total_rows += len(rows)
        made += 1

        # 첫 성공 시트를 미리보기로
        if preview is None:
            cols = ["출처파일", "출처폴더"]
            if header is not None:
                cols += [_fmt(h) for h in header]
            else:
                cols += [get_column_letter(i + 1) for i in range(n_cols)]
            prev_rows = []
            for file_base, src_folder, row in rows[:8]:
                padded = list(row) + [None] * (n_cols - len(row))
                prev_rows.append([file_base, src_folder] + [_fmt(v) for v in padded])
            preview = {"sheet": sheet_name, "columns": cols, "rows": prev_rows}

    # 대상 파일 로그 (취합된 파일 표시)
    included = set()
    for sr in sheet_results:
        # skipped 에 없으면 포함된 것으로 간주 (근사)
        pass
    for path in files:
        rel = os.path.relpath(path, folder)
        note, ok = file_notes.get(rel, ("취합 완료", True))
        all_log.append({"file": rel, "note": note, "ok": ok})

    if made == 0:
        return {"ok": False,
                "error": "취합된 시트가 없습니다. 시트명/헤더 행번호를 확인하세요.",
                "sheets": sheet_results, "log": all_log}

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_name = f"[취합]{prefix}_{ts}.xlsx"
    out_path = os.path.join(folder, out_name)
    out_wb.save(out_path)

    return {
        "ok": True,
        "out_path": out_path,
        "out_name": out_name,
        "made": made,
        "total_rows": total_rows,
        "file_count": len(files),
        "sheet_count": len(sheets),
        "sheets": sheet_results,
        "log": all_log,
        "preview": preview,
    }
