# -*- coding: utf-8 -*-
"""폴더 생성 기능 (create_folders.py 의 핵심 로직을 함수로 재사용).

원본 CLI/GUI 대신, 웹 UI 에서 넘어온 인자로 동작하도록 정리했다.
- 상위/하위 컬럼 값을 폴더 구조로 만든다.
- 폴더명 앞 번호는 '없음 / 2자리숫자(01.) / 알파벳(A.)' 중 선택.
"""

import re
from pathlib import Path

from openpyxl import load_workbook

# Windows 폴더 이름에 쓸 수 없는 문자
INVALID_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def sanitize(name) -> str:
    """폴더 이름으로 안전하게 정리한다."""
    name = str(name).strip()
    name = INVALID_CHARS.sub("_", name)
    name = name.rstrip(". ")  # Windows는 이름 끝의 점/공백을 허용하지 않음
    return name


def detect_header(ws):
    """헤더 행을 추정한다: 비어있지 않은 셀이 2개 이상인 첫 행.
    반환: (header_row_index, [(표시이름, 열인덱스), ...])  ※ 0-based row index
    못 찾으면 (None, [])."""
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        cells = [("" if c is None else str(c).strip()) for c in row]
        non_empty = [c for c in cells if c != ""]
        if len(non_empty) >= 2:
            headers = []
            for i, c in enumerate(cells):
                display = c if c != "" else f"(빈 열 {i + 1})"
                headers.append((display, i))
            return row_idx, headers
        if row_idx > 30:
            break
    return None, []


def get_workbook_info(excel_path: Path):
    """파일의 시트별 헤더 정보를 읽어온다.
    반환: {시트이름: (header_row_index, [(표시이름, 열인덱스), ...])}"""
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    info = {}
    for name in wb.sheetnames:
        ws = wb[name]
        info[name] = detect_header(ws)
    wb.close()
    return info


def collect_tree(excel_path: Path, sheet, upper_idx, header_row, lower_idx=None):
    """엑셀에서 상위/하위 값을 순서대로 수집한다.
    반환: dict{상위값: [하위값...]}  (하위가 없으면 값이 None)
    입력 순서(첫 등장 순서)를 유지하며 중복은 제거한다."""
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb[sheet]

    tree = {}
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx <= header_row:
            continue
        upper = row[upper_idx] if upper_idx < len(row) else None
        if upper is None or str(upper).strip() == "":
            continue
        u = sanitize(upper)
        if not u:
            continue

        if lower_idx is None:
            tree.setdefault(u, None)
            continue

        lower = row[lower_idx] if lower_idx < len(row) else None
        if lower is None or str(lower).strip() == "":
            continue
        l = sanitize(lower)
        if not l:
            continue
        lst = tree.setdefault(u, [])
        if lst is None:  # 같은 상위값이 하위 없이 먼저 등장했던 경우 방어
            lst = []
            tree[u] = lst
        if l not in lst:
            lst.append(l)

    wb.close()
    return tree


def _prefix(i, mode):
    """번호 접두어. mode: 'none' | 'digit' | 'alpha' (i 는 1-based)."""
    if mode == "digit":
        return f"{i:02d}."
    if mode == "alpha":
        # 1->A, 2->B ... 26->Z, 27->AA ...
        s = ""
        n = i
        while n > 0:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        return s + "."
    return ""


def build_preview(tree, number_mode="digit", top_limit=None):
    """UI 미리보기용 트리 구조를 만든다.
    반환: {'tree': [{'name':.., 'children':[{'name':..}]}], 'count_top':, 'count_total':}"""
    items = list(tree.items())
    count_top = len(items)
    count_total = sum(len(v) if v else 1 for v in tree.values())

    shown = items if top_limit is None else items[:top_limit]
    out = []
    for i, (upper, lowers) in enumerate(shown, start=1):
        u_name = _prefix(i, number_mode) + upper
        children = []
        if lowers:
            for j, lower in enumerate(lowers, start=1):
                children.append({"name": _prefix(j, number_mode) + lower})
        out.append({"name": u_name, "children": children})
    return {"tree": out, "count_top": count_top, "count_total": count_total}


def create_folders(tree, output_base: Path, number_mode="digit", dry_run=False, log=None):
    """폴더를 생성하고 (생성수, 기존수, 총폴더수, 로그) 를 반환한다."""
    created = existed = total = 0
    logs = []

    def _log(msg):
        logs.append(msg)
        if log:
            log(msg)

    def make(path):
        nonlocal created, existed
        if dry_run:
            _log(f"[미리보기] {path}")
            return
        if path.exists():
            existed += 1
            _log(f"[존재] {path}")
        else:
            path.mkdir(parents=True, exist_ok=True)
            created += 1
            _log(f"[생성] {path}")

    for i, (upper, lowers) in enumerate(tree.items(), start=1):
        u_name = _prefix(i, number_mode) + upper
        upper_path = output_base / u_name
        if not lowers:  # None 또는 빈 리스트 → 상위 폴더만 생성
            total += 1
            make(upper_path)
        else:
            for j, lower in enumerate(lowers, start=1):
                l_name = _prefix(j, number_mode) + lower
                total += 1
                make(upper_path / l_name)
    return created, existed, total, logs


# ── 웹 UI 용 상위 함수 ──────────────────────────────────────────
def load_excel(path):
    """엑셀의 시트/컬럼 목록을 UI 용으로 반환한다."""
    info = get_workbook_info(Path(path))
    sheets = []
    for name, (header_row, headers) in info.items():
        col_names = [disp for disp, _ in headers]
        sheets.append({
            "name": name,
            "header_row": header_row,
            "headers": col_names,
        })
    return {"ok": True, "sheets": sheets}


def _resolve_idx(path, sheet, col_name):
    info = get_workbook_info(Path(path))
    header_row, headers = info[sheet]
    name_to_idx = {name: idx for name, idx in headers}
    return header_row, name_to_idx.get(col_name)


def preview(path, sheet, upper_col, lower_col, number_mode="digit"):
    """생성될 폴더 구조 미리보기."""
    header_row, upper_idx = _resolve_idx(path, sheet, upper_col)
    if upper_idx is None:
        return {"ok": False, "error": f"상위 컬럼 '{upper_col}' 을(를) 찾을 수 없습니다."}
    lower_idx = None
    if lower_col and lower_col not in ("(없음)", ""):
        _, lower_idx = _resolve_idx(path, sheet, lower_col)
    tree = collect_tree(Path(path), sheet, upper_idx, header_row, lower_idx)
    if not tree:
        return {"ok": False, "error": "생성할 데이터가 없습니다. 컬럼/시트를 확인하세요."}
    result = build_preview(tree, number_mode)
    result["ok"] = True
    return result


def run(path, sheet, upper_col, lower_col, number_mode, output_dir):
    """실제 폴더 생성."""
    header_row, upper_idx = _resolve_idx(path, sheet, upper_col)
    if upper_idx is None:
        return {"ok": False, "error": f"상위 컬럼 '{upper_col}' 을(를) 찾을 수 없습니다."}
    lower_idx = None
    if lower_col and lower_col not in ("(없음)", ""):
        _, lower_idx = _resolve_idx(path, sheet, lower_col)

    tree = collect_tree(Path(path), sheet, upper_idx, header_row, lower_idx)
    if not tree:
        return {"ok": False, "error": "생성할 데이터가 없습니다."}

    created, existed, total, logs = create_folders(
        tree, Path(output_dir), number_mode=number_mode
    )
    return {
        "ok": True,
        "created": created,
        "existed": existed,
        "total": total,
        "output_dir": str(output_dir),
        "log": logs[-40:],  # 너무 길면 뒤쪽만
    }
